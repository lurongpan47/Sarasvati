#!/usr/bin/env python3
"""Build Microsoft Word (.docx) of the Jonang catalog Chinese translation."""
import pathlib, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = pathlib.Path(__file__).parent
raw = (ROOT / 'jonang-catalog_zh.md').read_text()

doc = Document()

# Set default font to CJK-safe
def set_cjk_font(run, font='SimSun', eastasia='SimSun', size=10):
    run.font.name = font
    run.font.size = Pt(size)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), eastasia)
    rFonts.set(qn('w:ascii'), font)
    rFonts.set(qn('w:hAnsi'), font)

# Document styles
style = doc.styles['Normal']
style.font.name = 'SimSun'
style.font.size = Pt(10.5)
style_rpr = style.element.get_or_add_rPr()
rFonts = OxmlElement('w:rFonts')
rFonts.set(qn('w:eastAsia'), 'SimSun')
rFonts.set(qn('w:ascii'), 'SimSun')
rFonts.set(qn('w:hAnsi'), 'SimSun')
style_rpr.append(rFonts)

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

# ---- Cover ----
cover_space = doc.add_paragraph()
cover_space.add_run('\n\n\n\n\n\n')

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_r = title_p.add_run('觉囊·大智慧藏总目录')
title_r.font.size = Pt(28)
title_r.font.bold = True
title_r.font.color.rgb = RGBColor(0x8B, 0x25, 0x00)
set_cjk_font(title_r, 'SimSun', 'SimSun', 28)
title_r.font.bold = True
title_r.font.color.rgb = RGBColor(0x8B, 0x25, 0x00)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_r = sub_p.add_run('ཇོ་ནང་རིག་མཛོད་ཆེན་མོའི་དཀར་ཆག')
sub_r.font.size = Pt(14)
sub_r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()

draft_p = doc.add_paragraph()
draft_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
draft_r = draft_p.add_run('汉译草稿 · AI 一稿 · 待复核')
draft_r.font.size = Pt(13)
draft_r.font.italic = True
draft_r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
set_cjk_font(draft_r, 'SimSun', 'SimSun', 13)
draft_r.font.italic = True
draft_r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

for _ in range(5): doc.add_paragraph()

# Metadata table on cover
cover_info = [
    ('原本', '觉囊文丛 · 第 216 册'),
    ('页数', '449 页藏文原本'),
    ('字符数', '约 99 万字藏文'),
    ('翻译', 'Claude Opus 4.7'),
    ('日期', '2026-10-17'),
    ('许可', 'CC BY-SA 4.0'),
    ('状态', 'AI 一稿 ⟨བརྟག⟩ · 所有译文待具名藏传学者复核'),
    ('归属', 'Sarasvatī 第八支（藏传佛典）档案项目'),
]
table = doc.add_table(rows=len(cover_info), cols=2)
table.autofit = False
for i, (k, v) in enumerate(cover_info):
    r = table.rows[i]
    r.cells[0].width = Cm(3.5)
    r.cells[1].width = Cm(12.5)
    pk = r.cells[0].paragraphs[0]
    rk = pk.add_run(k)
    rk.font.bold = True
    rk.font.color.rgb = RGBColor(0x8B, 0x25, 0x00)
    rk.font.size = Pt(11)
    set_cjk_font(rk, 'SimSun', 'SimSun', 11)
    rk.font.bold = True
    rk.font.color.rgb = RGBColor(0x8B, 0x25, 0x00)
    pv = r.cells[1].paragraphs[0]
    rv = pv.add_run(v)
    rv.font.size = Pt(11)
    set_cjk_font(rv, 'SimSun', 'SimSun', 11)

doc.add_page_break()

# ---- Body ----
sections = re.split(r'<!-- ============ Source PDF page (\d+) ============ -->', raw)
# sections[0] is the markdown header (which we replaced with the cover above)
pairs = list(zip(sections[1::2], sections[2::2]))

def add_source_marker(doc, pnum, note=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if note else WD_ALIGN_PARAGRAPH.CENTER
    text = f'— 源 PDF p.{pnum} —' if note else f'— 源 PDF 第 {pnum} 页（空白）—'
    r = p.add_run(text)
    r.font.size = Pt(8)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    set_cjk_font(r, 'SimSun', 'SimSun', 8)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

def add_h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x8B, 0x25, 0x00)
    set_cjk_font(r, 'SimSun', 'SimSun', 14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x8B, 0x25, 0x00)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)

def add_page_marker(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    set_cjk_font(r, 'SimSun', 'SimSun', 9)
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)

def add_body_line(doc, text):
    # Convert **bold** inline markers
    p = doc.add_paragraph()
    # Preserve leading full-width spaces visually
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            inner = part.strip('*')
            r = p.add_run(inner)
            r.font.size = Pt(10.5)
            r.font.bold = True
            set_cjk_font(r, 'SimSun', 'SimSun', 10.5)
            r.font.bold = True
        else:
            r = p.add_run(part)
            r.font.size = Pt(10.5)
            set_cjk_font(r, 'SimSun', 'SimSun', 10.5)
    p.paragraph_format.space_after = Pt(0)

for pnum, content in pairs:
    text = content.strip()
    if not text or '（空白页）' in text:
        add_source_marker(doc, pnum, note=False)
        continue
    for raw_line in text.split('\n'):
        line = raw_line.rstrip()
        if not line:
            continue
        if line.startswith('**') and line.endswith('**') and line.count('**') == 2:
            add_h1(doc, line.strip('*').strip())
        elif line.startswith('【') and '】' in line and len(line) < 20:
            add_page_marker(doc, line)
        else:
            add_body_line(doc, line)
    add_source_marker(doc, pnum, note=True)

# ---- Save ----
out = ROOT / 'jonang-catalog_zh.docx'
doc.save(str(out))
import os
print(f'built {out.name}: {os.path.getsize(out):,} bytes')
